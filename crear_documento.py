from pathlib import Path
import json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).parent
OUT = ROOT.parent / "Proyecto_Metodo_Simplex_Prototipo_1_0.docx"
EVIDENCE = json.loads((ROOT / "evidencia.json").read_text(encoding="utf-8"))
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.85)
sec.left_margin = sec.right_margin = Inches(.95)
sec.header_distance = sec.footer_distance = Inches(.38)
sec.different_first_page_header_footer = True
for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
    s = doc.styles[name]
    s.font.name = "Times New Roman"
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal = doc.styles["Normal"]
normal.font.size = Pt(11.5)
normal.paragraph_format.line_spacing = 1.1
normal.paragraph_format.space_after = Pt(6)
for name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11.5)]:
    s = doc.styles[name]
    s.font.size = Pt(size)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(10)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True
doc.styles["Title"].font.size = Pt(23)
doc.styles["Title"].paragraph_format.space_after = Pt(12)
doc.styles["Subtitle"].font.size = Pt(15)
doc.styles["Caption"].font.size = Pt(9.5)
doc.styles["Caption"].font.italic = True
doc.styles["Caption"].paragraph_format.space_after = Pt(8)
for name in ("List Bullet", "List Number"):
    doc.styles[name].font.name = "Times New Roman"
    doc.styles[name].font.size = Pt(11.5)
    doc.styles[name].paragraph_format.space_after = Pt(4)
settings = doc.settings.element
lang = OxmlElement("w:themeFontLang"); lang.set(qn("w:val"), "es-GT"); settings.append(lang)
for name in ("Normal", "Title", "Heading 1", "Heading 2"):
    lang = OxmlElement("w:lang"); lang.set(qn("w:val"), "es-GT")
    doc.styles[name].element.get_or_add_rPr().append(lang)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run()
r.font.size = Pt(10)
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
r._r.addnext(fld)

def p(text="", style=None, bold=False):
    para = doc.add_paragraph(style=style)
    rr = para.add_run(text); rr.bold = bold
    return para

def h(text, level=1):
    return doc.add_heading(text, level=level)

def page():
    doc.add_page_break()

def bullet(text):
    return p(text, "List Bullet")

def code(text, size=9.3):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.line_spacing = 1.03
    para.paragraph_format.keep_together = True
    rr = para.add_run(text)
    rr.font.name = "Consolas"; rr.font.size = Pt(size)
    return para

def table(headers, rows, widths, numeric=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i,w in enumerate(widths):
        t.columns[i].width = Inches(w)
    for i, label in enumerate(headers):
        t.rows[0].cells[i].text = str(label)
    for row in rows:
        for i,val in enumerate(row):
            if i == 0: cells = t.add_row().cells
            cells[i].text = str(val)
    for ri,row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        keep = OxmlElement("w:cantSplit"); pr.append(keep)
        if ri == 0:
            repeat = OxmlElement("w:tblHeader"); pr.append(repeat)
        for ci,c in enumerate(row.cells):
            c.width = Inches(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = c._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side,val in [("top","85"),("bottom","85"),("left","105"),("right","105")]:
                el = OxmlElement("w:"+side); el.set(qn("w:w"), val); el.set(qn("w:type"), "dxa"); margins.append(el)
            tcPr.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E4EAF0" if ri == 0 else ("F7F8F9" if ri % 2 == 0 else "FFFFFF"))
            tcPr.append(shading)
            borders=OxmlElement("w:tcBorders")
            for side in ["top","left","bottom","right"]:
                el = OxmlElement("w:"+side); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "D9D9D9"); borders.append(el)
            tcPr.append(borders)
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if numeric or ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.bold = ri == 0
    p().paragraph_format.space_after = Pt(1)
    return t

def math_run(text):
    rr=OxmlElement("m:r")
    props=OxmlElement("m:rPr")
    sty=OxmlElement("m:sty"); sty.set(qn("m:val"),"p"); props.append(sty); rr.append(props)
    tt=OxmlElement("m:t"); tt.text=text; rr.append(tt)
    return rr

def math_tokens(parent, text):
    for piece in re.split(r"([A-Za-z]+_[A-Za-z0-9]+|[A-Za-z]+\^[A-Za-z0-9]+)", text):
        if not piece: continue
        if "_" in piece and re.fullmatch(r"[A-Za-z]+_[A-Za-z0-9]+",piece):
            base,sub=piece.split("_"); node=OxmlElement("m:sSub")
            e=OxmlElement("m:e"); e.append(math_run(base)); s=OxmlElement("m:sub"); s.append(math_run(sub)); node.extend([e,s]); parent.append(node)
        elif "^" in piece and re.fullmatch(r"[A-Za-z]+\^[A-Za-z0-9]+",piece):
            base,sup=piece.split("^"); node=OxmlElement("m:sSup")
            e=OxmlElement("m:e"); e.append(math_run(base)); s=OxmlElement("m:sup"); s.append(math_run(sup)); node.extend([e,s]); parent.append(node)
        else: parent.append(math_run(piece))

def equation(text):
    para=doc.add_paragraph()
    para.paragraph_format.space_after=Pt(2)
    para.paragraph_format.line_spacing=1.0
    obj=OxmlElement("m:oMath"); math_tokens(obj,text); para._p.append(obj)
    return para

def link(label,url):
    para=p()
    hyper=OxmlElement("w:hyperlink")
    hyper.set(qn("r:id"),para.part.relate_to(url, RT.HYPERLINK,is_external=True))
    run=OxmlElement("w:r"); prop=OxmlElement("w:rPr")
    color=OxmlElement("w:color"); color.set(qn("w:val"),"1F4E79"); prop.append(color)
    run.append(prop); txt=OxmlElement("w:t"); txt.text=label; run.append(txt); hyper.append(run); para._p.append(hyper)
    return para

def capture(n):
    source=(ROOT/"capturas"/f"caso_{n}_tabla_final.txt").read_text(encoding="utf-8-sig")
    excerpt=source[source.index("TABLA FINAL"):].strip()
    code(excerpt,9.5)
    p(f"Salida {n}. Extracto de la ejecución real de simplex.py.", "Caption")
    p("Captura de pantalla de la ejecución pendiente de incorporar.").runs[0].italic=True

# Portada. Campos deliberadamente vacíos por solicitud del usuario.
p("Universidad  __________________________________________________", bold=True)
p("Facultad  _____________________________________________________", bold=True)
p().paragraph_format.space_after=Pt(14)
q=p("Investigación de Operaciones"); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
p().paragraph_format.space_after=Pt(17)
q=p("Calculadora del Método\nSimplex en Python", "Title"); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
q=p("Prototipo 1.0", "Subtitle"); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
p().paragraph_format.space_after=Pt(20)
p("Catedrático  __________________________________________________")
p().paragraph_format.space_after=Pt(10)
p("Integrantes del equipo",bold=True)
table(["Nombre completo","Carné"],[["",""],["",""],["",""],["",""]],[4.85,1.75])
p().paragraph_format.space_after=Pt(17)
p("Fecha de entrega  ______________________________________________")
page()

h("2 Introducción")
p("La programación lineal permite elegir valores para un conjunto de variables con el propósito de maximizar un beneficio o minimizar un costo, respetando restricciones de recursos. Se aplica, por ejemplo, a la planificación de producción, la asignación de materiales y la distribución de presupuestos. El Método Simplex busca una solución óptima mediante cambios de base que conservan la factibilidad del modelo.")
p("Este proyecto presenta un prototipo funcional del Simplex tabular en Python. El programa convierte restricciones de tipo menor o igual en ecuaciones, construye una tabla numérica y muestra cada pivoteo. Los dos ejemplos incorporados alcanzan los valores óptimos Z = 36 y Z = 21; la salida se contrasta con soluciones teóricas y con pruebas automatizadas.")
h("3 Objetivos")
h("Objetivo general",2)
p("Desarrollar un programa informático en Python capaz de resolver modelos de programación lineal mediante el Método Simplex tabular, dentro del alcance establecido para el Prototipo 1.0.")
h("Objetivos específicos",2)
bullet("Traducir la lógica matemática del pivoteo a estructuras de datos computacionales y operaciones elementales sobre filas.")
bullet("Validar los resultados del algoritmo con al menos dos ejercicios resueltos en clase, comparando sus tablas y soluciones con la resolución teórica.")
p("Confirmación pendiente: se utilizan los dos ejemplos actuales del programa; falta confirmar su correspondencia con los ejercicios resueltos en clase.")
h("4 Alcance y limitaciones del Prototipo 1.0")
p("El prototipo resuelve maximización con variables continuas no negativas, restricciones de tipo ≤ y lados derechos no negativos. Admite una o más variables y restricciones, agrega holguras y muestra las tablas inicial, intermedias y final. Permite utilizar ejemplos definidos en el código o ingresar coeficientes por consola.")
p("Se excluyen minimización, restricciones ≥ o =, lados derechos negativos, variables libres, variables artificiales, Gran M, Dos Fases e interfaz gráfica. Tampoco se resuelven modelos enteros o no lineales. La lectura de modelos desde archivos y el análisis de sensibilidad quedan para fases posteriores.")
p("Los cálculos usan punto flotante. El programa detecta falta de acotación, limita las iteraciones y detiene una base repetida, pero no implementa una regla general anticiclaje ni técnicas avanzadas para modelos mal escalados.")
page()

h("5 Requisitos técnicos y entorno de desarrollo")
p("El lenguaje utilizado es Python. La evidencia de este documento se ejecutó el 2 de octubre de 2026 en Windows 11 con las siguientes versiones. Las bibliotecas fijadas requieren Python 3.11 o posterior.")
table(["Componente","Versión","Función"],[
["Python","3.12.14","Control del algoritmo y entrada por consola"],
["NumPy","2.3.5","Arreglos numéricos y operaciones matriciales"],
["Pandas","3.0.1","Presentación de tablas en consola"]],[1.15,.85,4.6])
p("El programa también se había verificado con Python 3.11.6. Las bibliotecas argparse, fractions, pathlib y unittest pertenecen a la biblioteca estándar; se utilizan para las opciones de ejecución, entrada de fracciones, archivos y pruebas.")
h("Instalación y ejecución",2)
p("Desde la carpeta del programa, instalar las dependencias en el mismo intérprete con el que se ejecutará el archivo:")
code("python -m pip install numpy==2.3.5 pandas==3.0.1\npython simplex.py\npython simplex.py --interactivo\npython simplex.py --reporte simplex_resultado.txt\npython -m unittest -v test_simplex",9.7)
p("La ejecución sin opciones resuelve los dos ejemplos. El modo interactivo solicita el número de variables, las restricciones y sus coeficientes; admite decimales y fracciones separados por espacios. La opción --reporte guarda la salida en un archivo .txt; --info muestra el alcance y las versiones.")
h("Estructuras de datos principales",2)
p("Para n variables de decisión y m restricciones, tabla es un arreglo bidimensional de NumPy de tamaño (m + 1) × (n + m + 1). Las primeras m filas contienen las restricciones y la última representa la función objetivo.")
code("Restricciones:  [ A  | I | b ]\nFila objetivo:  [-c  | 0 | 0 ]",10.5)
p("A almacena los coeficientes del modelo; I es la identidad de las variables de holgura y b es el lado derecho. La columna explícita de Z se omite porque su coeficiente es 1 en la fila objetivo y 0 en las restantes.")
p("Las listas nombres_columnas y base relacionan las posiciones de la matriz con X1, X2, … y S1, S2, …. El resultado es un diccionario con Z y las variables de decisión. Pandas crea un DataFrame para imprimir la tabla; el pivoteo se calcula en NumPy.")
page()

h("6 Arquitectura y lógica del algoritmo")
h("Conversión a forma estándar",2)
p("A cada restricción se agrega una variable de holgura no negativa. Con b no negativo, X = 0 y S = b constituyen una solución básica factible inicial.")
equation("AX + S = b")
equation("Z − c^T X = 0")
p("La fila objetivo contiene los coeficientes negados. Mientras exista un coeficiente menor que −10⁻⁹, puede seleccionarse una variable de entrada.")
h("Pseudocódigo del ciclo Simplex",2)
code("""INICIO
  Leer c, A y b
  Validar dimensiones, números finitos y b >= 0
  Agregar holguras y construir [A | I | b; -c | 0 | 0]
  Inicializar la base con S1, ..., Sm y mostrar tabla

  MIENTRAS exista un coeficiente de Z < -1e-9
    Si se supera el límite o se repite una base: detener
    Elegir la columna con coeficiente más negativo
    Calcular b_i / a_iq solo donde a_iq > 1e-9
    Si no hay fila elegible: informar "no acotado"
    Elegir la fila con el cociente mínimo
    Mostrar entrada, salida, razones y elemento pivote
    Dividir toda la fila pivote por el elemento pivote
    Anular la columna pivote en las demás filas, incluida Z
    Actualizar la base, contar iteración y mostrar tabla
  FIN MIENTRAS

  Extraer X y Z y verificar factibilidad y objetivo
  Si la verificación falla: informar error numérico
  Mostrar tabla final y devolver la solución óptima
FIN""",9.5)
h("Criterios de selección",2)
p("La variable de entrada corresponde al coeficiente más negativo de la fila Z. La variable de salida se determina con el menor cociente entre el lado derecho y un coeficiente positivo de la columna pivote. El cociente cero es válido. En empates exactos se selecciona la primera columna o fila en el orden de la tabla.")
h("Operaciones de Gauss Jordan",2)
p("La fila pivote se divide entre su pivote para convertirlo en 1. Después se resta a cada otra fila el múltiplo de la fila normalizada que anula su entrada en esa columna. Las mismas operaciones se aplican a la fila Z y a la columna b.")
code("T[p, :] = T[p, :] / T[p, q]\nPara cada i distinto de p:\n    factor = T[i, q]\n    T[i, :] = T[i, :] - factor * T[p, :]",9.6)
page()

h("Funciones principales del programa",2)
p("La clase SimplexTabular concentra los datos y las operaciones matemáticas. Las funciones de consola preparan los modelos, presentan la solución y conservan la evidencia de ejecución.")
table(["Función o método","Responsabilidad"],[
["_validar_datos()","Comprueba dimensiones, coeficientes finitos y b no negativo."],
["_construir_tabla_inicial()","Agrega la identidad de holguras y la fila objetivo."],
["es_optima()","Revisa la ausencia de coeficientes negativos significativos en Z."],
["encontrar_columna_pivote()","Selecciona el coeficiente más negativo de Z."],
["calcular_cocientes()","Obtiene razones y excluye denominadores no positivos."],
["encontrar_fila_pivote()","Selecciona la razón mínima o informa falta de acotación."],
["actualizar_tabla()","Normaliza la fila pivote y ejecuta Gauss-Jordan."],
["resolver()","Controla el ciclo, las paradas y la verificación final."],
["obtener_solucion()","Extrae Z y las variables; las no básicas valen cero."],
["verificar_solucion()","Comprueba X ≥ 0, AX ≤ b y la consistencia del objetivo."],
["resolver_caso()","Muestra el modelo y contrasta el resultado con la referencia."],
["modo_interactivo() y main()","Gestionan la entrada, las opciones y el reporte de texto."]],[2.75,3.85])
h("Control de errores y terminación",2)
p("Una fila con coeficiente cero o negativo en la columna pivote se excluye de la razón mínima. El método de actualización también rechaza un pivote nulo o demasiado pequeño. Si no hay fila elegible, se informa que el problema es no acotado dentro de la tolerancia usada.")
p("El límite predeterminado es de 1000 iteraciones por modelo y puede cambiarse con --max-iteraciones. El programa guarda las bases visitadas para detectar repeticiones y detener un posible ciclo. Estas paradas no se presentan como soluciones óptimas.")
p("La tolerancia de selección y optimalidad es 10⁻⁹. La comprobación final utiliza 10⁻⁷, con ajuste por escala para las restricciones y tolerancias absoluta y relativa para el objetivo. La salida muestra hasta ocho cifras significativas; las operaciones conservan los valores en punto flotante.")
page()

h("7 Pruebas y validación")
p("Se ejecutaron los dos ejemplos incorporados en simplex.py. Cada caso muestra la entrada, la tabla final y la comparación con una solución teórica. Está pendiente confirmar que correspondan a los ejercicios de clase.")
h("Caso de prueba 1",2)
equation("Maximizar Z = 3X_1 + 5X_2")
equation("X_1 ≤ 4")
equation("2X_2 ≤ 12")
equation("3X_1 + 2X_2 ≤ 18")
equation("X_1, X_2 ≥ 0")
p("Entrada definida en CASOS_PRUEBA y ejecutada con python simplex.py:")
code("c = [3, 5]\nA = [[1, 0], [0, 2], [3, 2]]\nb = [4, 12, 18]",9.5)
p("En la primera iteración entra X2 y sale S2, con pivote 2. En la segunda entra X1 y sale S3, con pivote 3.")
capture(1)
p("Interpretación. La solución es X1 = 2, X2 = 6 y Z = 36. Las holguras son S1 = 2, S2 = 0 y S3 = 0; las restricciones segunda y tercera se cumplen con igualdad.")
p("Comprobación teórica. Se verifica 2 ≤ 4, 2(6) = 12 y 3(2) + 2(6) = 18. Al multiplicar la segunda restricción por 1.5 y sumar la tercera se obtiene 3X1 + 5X2 ≤ 36. La solución alcanza esa cota, por lo que es óptima.")
page()

h("Caso de prueba 2",2)
equation("Maximizar Z = 5X_1 + 4X_2")
equation("6X_1 + 4X_2 ≤ 24")
equation("X_1 + 2X_2 ≤ 6")
equation("X_1, X_2 ≥ 0")
p("Entrada definida en CASOS_PRUEBA y ejecutada con python simplex.py:")
code("c = [5, 4]\nA = [[6, 4], [1, 2]]\nb = [24, 6]",9.5)
p("En la primera iteración entra X1 y sale S1, con pivote 6. En la segunda entra X2 y sale S2, con pivote 4/3.")
capture(2)
p("Interpretación. La solución es X1 = 3, X2 = 1.5 y Z = 21. Ambas holguras valen cero, por lo que las dos restricciones se cumplen con igualdad.")
p("Comprobación teórica. Se verifica 6(3) + 4(1.5) = 24 y 3 + 2(1.5) = 6. Multiplicar la primera restricción por 0.75 y la segunda por 0.5, y sumarlas, da 5X1 + 4X2 ≤ 21. La solución alcanza esa cota y es óptima.")
p("El resultado también coincide directamente con la función objetivo: 5(3) + 4(1.5) = 21. Las dos variables son no negativas.")
page()

h("Comparación de resultados y pruebas adicionales",2)
table(["Caso","X1","X2","Z teórico","Z obtenido","Iteraciones"],
[["1","2","6","36","36","2"],["2","3","1.5","21","21","2"]],[.55,.6,.6,1.3,1.3,1.45],numeric=True)
p("Los dos casos aprobaron los controles de no negatividad, restricciones, valor de la función objetivo y coincidencia con la referencia. Las cotas obtenidas como combinaciones no negativas de las restricciones proporcionan una comprobación de optimalidad independiente de las tablas del algoritmo.")
p("También se ejecutó test_simplex.py con unittest. Las 24 pruebas terminaron correctamente; incluyen la comparación de 20 modelos pequeños con enumeración independiente de vértices, además de los siguientes comportamientos:")
bullet("Rechazo de dimensiones inconsistentes, datos no numéricos, NaN, infinito y lados derechos negativos.")
bullet("Cocientes cero, empates, restricciones redundantes y soluciones óptimas en la tabla inicial.")
bullet("Detección de problemas no acotados, límite de iteraciones y repetición de una base en un caso de ciclo.")
bullet("Entrada interactiva con fracciones y decimales, reintentos de entrada y exportación de reportes.")
p("La evidencia conserva los modelos, las tablas y la salida de las pruebas. Estos resultados verifican el alcance implementado; la validación específica con ejercicios de clase queda pendiente de su identificación.")
h("8 Conclusiones de la Fase 1")
p("El prototipo traduce el pivoteo del Simplex a operaciones sobre un arreglo bidimensional. La identidad de las holguras permite iniciar con una base factible para el alcance admitido, y las listas de nombres mantienen la relación entre los números de la tabla y las variables del modelo.")
p("El programa muestra cada iteración y obtiene los resultados esperados de los dos ejemplos. La sustitución en las restricciones y las cotas superiores confirman Z = 36 y Z = 21. Las pruebas automatizadas aportan evidencia adicional sobre los controles de entrada y las condiciones de terminación.")
p("Los desafíos técnicos del diseño fueron preservar la fila pivote normalizada, actualizar correctamente la base, excluir divisiones por cero y evitar ciclos indefinidos. Se atendieron con operaciones sobre una copia de la tabla, validaciones del pivote, tolerancias numéricas y controles de bases e iteraciones.")
p("Se cumple el objetivo de implementar el Método Simplex tabular dentro de las limitaciones declaradas. Para dar por cumplido el objetivo de validación con ejercicios resueltos en clase, aún debe confirmarse que los modelos documentados sean los utilizados por el catedrático.")
page()

h("9 Trabajo futuro para la Fase 2")
p("Las siguientes mejoras ampliarían el conjunto de modelos y facilitarían el uso del programa:")
bullet("Leer coeficientes desde archivos .txt o .csv con un formato definido y validación de errores.")
bullet("Aplicar una regla anticiclaje, como Bland, y ampliar el tratamiento de degeneración y empates.")
bullet("Implementar Dos Fases o Gran M para construir una base inicial cuando las holguras no sean suficientes.")
bullet("Incorporar minimización, restricciones ≥ o =, lados derechos negativos y transformaciones de variables.")
bullet("Mejorar el escalado numérico y detectar soluciones óptimas alternativas.")
bullet("Desarrollar una interfaz gráfica para introducir el modelo y consultar sus iteraciones.")
h("10 Referencias y bibliografía")
p("Cornell University, CS 6820. (2020, octubre). Lecture notes on the simplex method. Apuntes académicos sobre bases, pivoteo y el algoritmo Simplex.")
link("Consultar los apuntes de Cornell", "https://www.cs.cornell.edu/courses/cs6820/2021fa/handouts/splex.pdf")
p("Python Software Foundation. (s. f.). Documentación de Python 3.11. Biblioteca estándar, lenguaje y herramientas de ejecución.")
link("Documentación oficial de Python 3.11", "https://docs.python.org/3.11/")
p("NumPy Developers. (s. f.). numpy.ndarray. NumPy v2.3 Manual. Representación de arreglos numéricos multidimensionales.")
link("Referencia oficial de numpy.ndarray", "https://numpy.org/doc/2.3/reference/generated/numpy.ndarray.html")
p("pandas development team. (s. f.). pandas.DataFrame. Referencia de la estructura tabular utilizada para presentar los resultados.")
link("Referencia oficial de pandas.DataFrame", "https://pandas.pydata.org/docs/reference/api/pandas.DataFrame.html")
p("Fuentes técnicas consultadas el 2 de octubre de 2026. La versión instalada de cada biblioteca se registra en el apartado 5.")
p("Material de clase para completar la validación")
p("Referencia  ___________________________________________________")
p("________________________________________________________________")
doc.core_properties.title = "Calculadora del Método Simplex en Python"
doc.core_properties.subject = "Prototipo 1.0 de Investigación de Operaciones"
doc.core_properties.author = ""
doc.core_properties.keywords = "Simplex, programación lineal, Python"
doc.core_properties.comments = ""
doc.save(OUT)
print(OUT)
